from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

root = Path(__file__).parent
source = root / "guia-github-actions-e-seguranca.md"
target = root / "guia-github-actions-e-seguranca.docx"
lines = source.read_text(encoding="utf-8").splitlines()
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.68)
sec.bottom_margin = Inches(.68)
sec.left_margin = Inches(.78)
sec.right_margin = Inches(.78)
styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(5)
for sty, size, color in [("Title", 23, "14365D"), ("Heading 1", 16, "14365D"), ("Heading 2", 12, "0070C0")]:
    styles[sty].font.name = "Aptos Display"
    styles[sty].font.size = Pt(size)
    styles[sty].font.bold = True
    styles[sty].font.color.rgb = RGBColor.from_string(color)

header = sec.header.paragraphs[0]
header.text = "SOM & CORDA  |  GUIA OPERACIONAL"
header.style = "Caption"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer = sec.footer.paragraphs[0]
footer.text = "Trilha de aprendizado • Atualizado em 10/10/2026"
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

def add_runs(p, text):
    # Keep inline code and bold text legible in Word; markdown links remain readable URLs.
    pattern = r"(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^\)]+\))"
    for part in re.split(pattern, text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string("5B2C83")
        elif part.startswith("[") and "](" in part:
            label, url = part[1:].split("](", 1)
            p.add_run(f"{label} — {url[:-1]}")
        else:
            p.add_run(part)

i = 0
first_title = True
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1
        continue
    if line.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", c or "-") for c in cells):
                rows.append(cells)
            i += 1
        if rows:
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Light Shading Accent 1"
            for ri, row in enumerate(rows):
                cells = table.add_row().cells
                for ci, value in enumerate(row):
                    cells[ci].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = cells[ci].paragraphs[0]
                    add_runs(p, value)
                    if ri == 0:
                        for run in p.runs:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
            doc.add_paragraph()
        continue
    image_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
    if image_match:
        alt, image_path = image_match.groups()
        image_para = doc.add_paragraph()
        image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_para.add_run().add_picture(str(root / image_path), width=Inches(6.35))
        caption = doc.add_paragraph(alt, style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1
        continue
    if line.startswith("# "):
        title = line[2:]
        p = doc.add_paragraph(style="Title" if first_title else "Heading 1")
        add_runs(p, title)
        first_title = False
    elif line.startswith("## "):
        p = doc.add_paragraph(style="Heading 1")
        add_runs(p, line[3:])
    elif line.startswith("### "):
        p = doc.add_paragraph(style="Heading 2")
        add_runs(p, line[4:])
    elif line.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.18)
        p.paragraph_format.right_indent = Inches(.18)
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(8)
        pPr = p._p.get_or_add_pPr()
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "EAF2F8")
        pPr.append(shd)
        add_runs(p, line[2:])
    elif re.match(r"\d+\. ", line):
        p = doc.add_paragraph(style="List Number")
        add_runs(p, re.sub(r"^\d+\. ", "", line))
    elif line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, line[2:])
    elif line.startswith("**") and line.endswith("**"):
        p = doc.add_paragraph()
        add_runs(p, line)
    else:
        p = doc.add_paragraph()
        add_runs(p, line)
    i += 1

doc.core_properties.title = "GitHub Actions e segurança do repositório"
doc.core_properties.subject = "Passos para administradores e participantes de PRs"
doc.core_properties.author = "Som & Corda"
doc.save(target)
print(target)
